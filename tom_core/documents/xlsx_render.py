"""Professional Excel workbooks: styled tables with real formulas, formats, filters, charts, a summary."""
from __future__ import annotations

import re
from datetime import date, datetime
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from tom_core.documents.design import Palette
from tom_core.documents.spec import DocSpec

FORMATS = {"number": "#,##0.##", "currency": "$#,##0.00", "percent": "0.0%", "date": "yyyy-mm-dd", "text": "@"}
_THIN = Side(style="thin", color="D1D5DB")


def _coerce(value: Any, kind: str) -> Any:
    if kind in ("number", "currency", "percent"):
        if isinstance(value, (int, float)):
            return value
        s = str(value).replace(",", "").replace("$", "").strip()
        pct = s.endswith("%")
        try:
            v = float(s.rstrip("%"))
            return v / 100 if pct or (kind == "percent" and v > 1) else v
        except ValueError:
            return value
    if kind == "date":
        try:
            return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
        except ValueError:
            return value
    return value


def _print_setup(ws) -> None:
    """Print/PDF: landscape, everything one page wide."""
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def _safe_name(name: str, used: set) -> str:
    n = "".join(ch for ch in str(name) if ch not in '[]:*?/\\')[:28] or "Sheet"
    base, k = n, 2
    while n in used:
        n = f"{base[:25]} {k}"
        k += 1
    used.add(n)
    return n


def render_workbook(content: Dict[str, Any], spec: DocSpec, pal: Palette, path: str) -> Dict[str, Any]:
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    head_fill = PatternFill("solid", fgColor=pal.primary)
    summary_rows: List[tuple] = []
    for idx, sheet in enumerate(content["sheets"][:5]):
        cols = [c if isinstance(c, dict) else {"name": str(c), "type": "text"} for c in sheet["columns"]]
        kinds = [c.get("type", "text") if c.get("type") in FORMATS else "text" for c in cols]
        rows = [list(r)[:len(cols)] + [""] * (len(cols) - len(r)) for r in sheet["rows"] if isinstance(r, list)][:500]
        # totals are added as live formulas below — drop any the model wrote itself
        rows = [r for r in rows if not re.match(r"^\s*(grand\s+)?(total|sum|average|avg)\b", str(r[0]), re.I)]
        ws = wb.create_sheet(_safe_name(sheet.get("name") or f"Sheet {idx + 1}", used))
        ws.sheet_view.showGridLines = False
        last_col = get_column_letter(len(cols))

        ws["A1"] = content["title"] if idx == 0 else (sheet.get("name") or "")
        ws["A1"].font = Font(name="Segoe UI Semibold", size=18, bold=True, color=pal.primary)
        ws["A2"] = sheet.get("description", "")
        ws["A2"].font = Font(name="Segoe UI", size=10, italic=True, color=pal.muted)
        # no merged title cells: long titles simply flow across the empty cells to the right
        ws.row_dimensions[1].height = 30

        hr = 4
        for c, col in enumerate(cols, 1):
            cell = ws.cell(row=hr, column=c, value=str(col["name"]))
            cell.font = Font(name="Segoe UI", bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=_THIN)
        ws.row_dimensions[hr].height = 26
        for r, row in enumerate(rows, hr + 1):
            for c, (val, kind) in enumerate(zip(row, kinds), 1):
                cell = ws.cell(row=r, column=c, value=_coerce(val, kind))
                cell.number_format = FORMATS[kind]
                cell.font = Font(name="Segoe UI", size=10.5)
                cell.border = Border(bottom=_THIN)
                if kind in ("number", "currency", "percent"):
                    cell.alignment = Alignment(horizontal="right")
        first, last = hr + 1, hr + len(rows)

        # Excel table (banded rows + filter buttons) over the data
        if rows:
            tbl = Table(displayName=f"T{idx + 1}_{''.join(ch for ch in ws.title if ch.isalnum())[:20] or 'Data'}",
                        ref=f"A{hr}:{last_col}{last}")
            tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(tbl)

        # Totals / averages with live formulas
        numeric = [i for i, k in enumerate(kinds, 1) if k in ("number", "currency")]
        if rows and numeric:
            tr = last + 1
            ws.cell(row=tr, column=1, value="Total").font = Font(bold=True, color=pal.primary)
            ws.cell(row=tr + 1, column=1, value="Average").font = Font(bold=True, color=pal.primary)
            for c in numeric:
                L = get_column_letter(c)
                for rr, fn in ((tr, "SUM"), (tr + 1, "AVERAGE")):
                    cell = ws.cell(row=rr, column=c, value=f"={fn}({L}{first}:{L}{last})")
                    cell.number_format = FORMATS[kinds[c - 1]]
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill("solid", fgColor=pal.light)
                    cell.border = Border(top=Side(style="medium", color=pal.primary))
                ws.conditional_formatting.add(f"{L}{first}:{L}{last}", ColorScaleRule(
                    start_type="min", start_color="FFFFFF", end_type="max", end_color=pal.light))
                summary_rows.append((ws.title, str(cols[c - 1]["name"]), f"'{ws.title}'!{L}{first}:{L}{last}",
                                     FORMATS[kinds[c - 1]]))

        # Column widths, freeze header
        for c, col in enumerate(cols, 1):
            longest = max([len(str(col["name"]))] + [len(str(r[c - 1])) for r in rows[:100]])
            ws.column_dimensions[get_column_letter(c)].width = min(max(12, longest + 4), 48)
        ws.freeze_panes = ws.cell(row=hr + 1, column=1)
        _print_setup(ws)

        # Chart
        ch = sheet.get("chart") or {}
        ctype = str(ch.get("type", "bar")).lower()
        names = [str(c["name"]) for c in cols]
        label_col = names.index(ch["label_column"]) + 1 if ch.get("label_column") in names else 1
        values = [names.index(v) + 1 for v in (ch.get("value_columns") or []) if v in names and
                  kinds[names.index(v)] in ("number", "currency", "percent")] or numeric[:2]
        if rows and values and ctype != "none":
            chart = {"line": LineChart, "pie": PieChart}.get(ctype, BarChart)()
            if ctype != "pie":
                chart.varyColors = False
            chart.title = ch.get("title") or f"{names[values[0] - 1]} by {names[label_col - 1]}"
            chart.height, chart.width = 8.5, 17
            for v in (values[:1] if ctype == "pie" else values[:3]):
                chart.add_data(Reference(ws, min_col=v, min_row=hr, max_row=min(last, hr + 30)), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=label_col, min_row=first, max_row=min(last, hr + 30)))
            if ctype != "pie":
                chart.y_axis.majorGridlines = None
                palette = [pal.primary, pal.secondary, pal.accent]
                for k, ser in enumerate(chart.series):
                    ser.graphicalProperties.solidFill = palette[k % 3]
                    ser.graphicalProperties.line.solidFill = palette[k % 3]
                if len(chart.series) == 1:
                    chart.legend = None
                else:
                    chart.legend.position = "b"
            ws.add_chart(chart, f"{get_column_letter(len(cols) + 2)}{hr}")

        if not content.get("data_from_user"):
            note = ws.cell(row=(last if rows else hr) + 4, column=1,
                           value="Note: sample figures generated by TOM for illustration — replace with your own data."
                           if not content.get("data_from_research") else "Figures based on web research; verify sources.")
            note.font = Font(italic=True, size=9, color=pal.muted)

    # Summary sheet with formulas over every numeric column
    if summary_rows:
        ws = wb.create_sheet("Summary", 0)
        ws.sheet_view.showGridLines = False
        ws["A1"] = f"{content['title']} — Summary"
        ws["A1"].font = Font(name="Segoe UI Semibold", size=18, bold=True, color=pal.primary)
        ws["A2"] = f"Prepared {date.today():%B %d, %Y}" + (f" by {spec.sender}" if spec.sender else "")
        ws["A2"].font = Font(italic=True, size=10, color=pal.muted)
        heads = ["Sheet", "Measure", "Total", "Average", "Minimum", "Maximum"]
        for c, h in enumerate(heads, 1):
            cell = ws.cell(row=4, column=c, value=h)
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head_fill
        for r, (sheet_name, measure, ref, fmt) in enumerate(summary_rows, 5):
            ws.cell(row=r, column=1, value=sheet_name)
            ws.cell(row=r, column=2, value=measure)
            for c, fn in enumerate(("SUM", "AVERAGE", "MIN", "MAX"), 3):
                cell = ws.cell(row=r, column=c, value=f"={fn}({ref})")
                cell.number_format = fmt
        for c, wdt in zip("ABCDEF", (22, 26, 16, 16, 16, 16)):
            ws.column_dimensions[c].width = wdt
        _print_setup(ws)
    wb.properties.title = content["title"]
    wb.properties.creator = spec.sender or "TOM"
    wb.save(path)
    return {"path": path, "sheets": len(wb.sheetnames)}
