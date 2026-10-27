"""Professional Word documents: cover page, contents, styled headings, tables, header/footer, page numbers.
Letters get a business-letter layout instead (no cover page)."""
from __future__ import annotations

from datetime import date
from typing import Any, Dict, List

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

from tom_core.documents.design import Palette, rgb, split_inline
from tom_core.documents.spec import DocSpec


def _color(hex_color: str) -> RGBColor:
    return RGBColor(*rgb(hex_color))


def _shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _field(paragraph, instr: str) -> None:
    """Insert a Word field (PAGE, NUMPAGES) — Word fills it in when the page renders."""
    # One run per part, exactly as Word writes fields (all parts in one run made Word hang on PDF export).
    def char(kind: str) -> None:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        paragraph.add_run()._r.append(el)

    char("begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    paragraph.add_run()._r.append(it)
    char("separate")
    paragraph.add_run("1")
    char("end")


def _bottom_border(paragraph, hex_color: str, size: int = 12) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(size)), ("w:space", "4"), ("w:color", hex_color)):
        b.set(qn(k), v)
    borders.append(b)
    p_pr.append(borders)


def _runs(paragraph, text: str, size: float = None, color: str = None, bold: bool = False) -> None:
    for part, is_bold, is_code in split_inline(text):
        r = paragraph.add_run(part)
        r.bold = bold or is_bold
        if is_code:
            r.font.name = "Consolas"
            r.font.color.rgb = _color("9D174D")
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = _color(color)


def _code(doc, text: str, pal: Palette) -> None:
    t = doc.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    _shade(c, "F3F4F6")
    c.text = ""
    for i, line in enumerate(text.split("\n")):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line or " ")
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        r.font.color.rgb = _color("111827")
    doc.add_paragraph()


def _styles(doc, pal: Palette) -> None:
    st = doc.styles
    normal = st["Normal"]
    normal.font.name = pal.body_font
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), pal.body_font)
    normal.font.size = Pt(11)
    normal.font.color.rgb = _color(pal.text)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color, before in (("Heading 1", 17, pal.primary, 18), ("Heading 2", 13, pal.secondary, 12),
                                      ("Title", 30, pal.primary, 0)):
        s = st[name]
        s.font.name = pal.heading_font
        s.element.rPr.rFonts.set(qn("w:eastAsia"), pal.heading_font)
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = _color(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True


def _header_footer(section, title: str, pal: Palette) -> None:
    hp = section.header.paragraphs[0]
    hp.text = ""
    _runs(hp, title, size=8.5, color=pal.muted)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Page ")
    r.font.size, r.font.color.rgb = Pt(8.5), _color(pal.muted)
    _field(fp, "PAGE")
    r = fp.add_run(" of ")
    r.font.size, r.font.color.rgb = Pt(8.5), _color(pal.muted)
    _field(fp, "NUMPAGES")
    for run in fp.runs:                                   # field results take the footer's size and colour too
        run.font.size, run.font.color.rgb = Pt(8.5), _color(pal.muted)


def _table(doc, block: Dict[str, Any], pal: Palette) -> None:
    headers, rows = block["headers"], block["rows"]
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        _runs(c.paragraphs[0], str(h), size=10, color="FFFFFF", bold=True)
        _shade(c, pal.primary)
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            _runs(cells[i].paragraphs[0], str(v), size=10)
            if n % 2:
                _shade(cells[i], pal.light)
    doc.add_paragraph()


def _callout(doc, text: str, pal: Palette) -> None:
    t = doc.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    _shade(c, pal.light)
    c.text = ""
    p = c.paragraphs[0]
    _runs(p, "Key point: ", size=10.5, color=pal.primary, bold=True)
    _runs(p, text, size=10.5)
    doc.add_paragraph()


def _blocks(doc, blocks: List[Dict[str, Any]], pal: Palette) -> None:
    for b in blocks:
        kind = b["type"]
        if kind == "heading":
            doc.add_heading(b["text"], level=2)
        elif kind == "paragraph":
            p = doc.add_paragraph()
            _runs(p, b["text"])
        elif kind in ("bullets", "numbered"):
            for it in b["items"]:
                p = doc.add_paragraph(style="List Bullet" if kind == "bullets" else "List Number")
                if ":" in it and len(it.split(":", 1)[0]) < 45 and "**" not in it:
                    head, rest = it.split(":", 1)
                    _runs(p, head + ":", bold=True)
                    _runs(p, rest)
                else:
                    _runs(p, it)
        elif kind == "table":
            _table(doc, b, pal)
        elif kind == "callout":
            _callout(doc, b["text"], pal)
        elif kind == "code":
            _code(doc, b["text"], pal)


def render_document(content: Dict[str, Any], spec: DocSpec, pal: Palette, path: str) -> Dict[str, Any]:
    doc = Document()
    _styles(doc, pal)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    if content.get("letter"):
        return _letter(doc, content, spec, pal, path)
    if 0 < spec.pages <= 2:
        return _compact(doc, content, spec, pal, path)

    # Cover page
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    _runs(p, spec.genre.upper(), size=11, color=pal.accent, bold=True)
    title = doc.add_paragraph(style="Title")
    title.add_run(content["title"])
    _bottom_border(title, pal.accent, 18)
    if content.get("subtitle"):
        p = doc.add_paragraph()
        _runs(p, content["subtitle"], size=14, color=pal.secondary)
    for _ in range(8):
        doc.add_paragraph()
    meta = [f"Prepared by {spec.sender}" if spec.sender else "", f"Prepared for {spec.audience.title()}" if spec.audience else "",
            f"{date.today():%B %d, %Y}"]
    for line in [m for m in meta if m]:
        p = doc.add_paragraph()
        _runs(p, line, size=11, color=pal.muted)

    # Contents + body in a new section with header/footer
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.header.is_linked_to_previous = False                   # the cover keeps an empty header/footer
    body.footer.is_linked_to_previous = False
    _header_footer(body, content["title"], pal)
    doc.add_heading("Contents", level=1)
    for i, s in enumerate(content["sections"], 1):
        p = doc.add_paragraph()
        _runs(p, f"{i}.  {s['heading']}", size=11)
        p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    for i, s in enumerate(content["sections"], 1):
        doc.add_heading(f"{i}. {s['heading']}", level=1)
        _blocks(doc, s["blocks"], pal)
    doc.core_properties.title = content["title"]
    doc.core_properties.author = spec.sender or "TOM"
    doc.core_properties.subject = spec.topic
    doc.save(path)
    return {"path": path, "sections": len(content["sections"])}


def _compact(doc, content, spec: DocSpec, pal: Palette, path: str) -> Dict[str, Any]:
    """One- or two-page documents: a title block instead of a cover page and contents page ("a one-page
    proposal" came out as cover + contents + 1,749 words in a live check, 2026-10-27)."""
    _header_footer(doc.sections[0], content["title"], pal)
    p = doc.add_paragraph()
    _runs(p, spec.genre.upper(), size=10, color=pal.accent, bold=True)
    title = doc.add_paragraph(style="Title")
    title.add_run(content["title"])
    _bottom_border(title, pal.accent, 12)
    meta = " · ".join(m for m in (content.get("subtitle") or "", f"Prepared by {spec.sender}" if spec.sender else "",
                                  f"{date.today():%B %d, %Y}") if m)
    p = doc.add_paragraph()
    _runs(p, meta, size=10, color=pal.muted)
    for s in content["sections"]:
        doc.add_heading(s["heading"], level=2)
        _blocks(doc, s["blocks"], pal)
    doc.core_properties.title = content["title"]
    doc.core_properties.author = spec.sender or "TOM"
    doc.core_properties.subject = spec.topic
    doc.save(path)
    return {"path": path, "sections": len(content["sections"])}


def _letter(doc, content, spec: DocSpec, pal: Palette, path: str) -> Dict[str, Any]:
    if spec.sender:
        p = doc.add_paragraph()
        _runs(p, spec.sender, size=14, color=pal.primary, bold=True)
        _bottom_border(p, pal.accent, 10)
    p = doc.add_paragraph()
    _runs(p, f"{date.today():%B %d, %Y}", color=pal.muted)
    doc.add_paragraph()
    to = (spec.recipient or "").strip()
    if to:
        p = doc.add_paragraph()
        _runs(p, "To:", bold=True)
        p = doc.add_paragraph()
        _runs(p, to[:1].upper() + to[1:])
    doc.add_paragraph()
    p = doc.add_paragraph()
    _runs(p, "Subject: ", bold=True)
    _runs(p, content["subject"])
    doc.add_paragraph()
    greet = f"Dear {to[:1].upper() + to[1:]}," if to and len(to.split()) <= 4 else "Dear Sir or Madam,"
    doc.add_paragraph(greet)
    _blocks(doc, content["sections"][0]["blocks"], pal)
    doc.add_paragraph()
    doc.add_paragraph("Sincerely,")
    doc.add_paragraph()
    p = doc.add_paragraph()
    _runs(p, spec.sender or "", bold=True)
    for para in doc.paragraphs:
        if para.style.name == "Normal" and para.text and not para.text.startswith(("Dear", "Sincerely", "To:", "Subject")):
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if len(para.text) > 120 else para.alignment
    doc.core_properties.title = content["subject"]
    doc.core_properties.author = spec.sender or "TOM"
    doc.save(path)
    return {"path": path, "sections": 1}
