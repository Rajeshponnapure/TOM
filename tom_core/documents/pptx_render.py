"""Professional, animated PowerPoint decks (16:9).

Every content slide has a designed layout — numbered points, cards, big-number stats, a native chart,
a timeline, a table, a quote or two columns — chosen from what the content actually is. Slides get a
fade transition and their elements fade in one after another (real PowerPoint entrance animations,
written as <p:timing> XML because python-pptx has no animation API). Speaker notes are included.
"""
from __future__ import annotations

import re
from datetime import date
from typing import Any, Dict, List, Optional, Tuple

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from tom_core.documents.design import Palette, rgb
from tom_core.documents.spec import DocSpec

W, H = Inches(13.333), Inches(7.5)
_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _c(hex_color: str) -> RGBColor:
    return RGBColor(*rgb(hex_color))


def _bg(slide, hex_color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _c(hex_color)


def _box(slide, x, y, w, h, color: str, shape=MSO_SHAPE.RECTANGLE, line: Optional[str] = None):
    s = slide.shapes.add_shape(shape, int(x), int(y), int(w), int(h))
    s.fill.solid()
    s.fill.fore_color.rgb = _c(color)
    if line:
        s.line.color.rgb = _c(line)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def _text(slide, x, y, w, h, text: str, size: float, color: str, *, bold=False, font="Segoe UI",
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, shape=None):
    s = shape or slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        for part in re.split(r"(\*\*[^*]+\*\*)", line):
            if not part:
                continue
            r = p.add_run()
            r.text = part.strip("*") if part.startswith("**") else part
            r.font.size = Pt(size)
            r.font.bold = bold or part.startswith("**")
            r.font.color.rgb = _c(color)
            r.font.name = font
    return s


def _fit(size: float, text: str, limit: int) -> float:
    """Shrink a font a little for long text so nothing overflows its box."""
    n = len(text)
    return size if n <= limit else max(size * 0.72, size * limit / n)


# ── animation + transition XML ──────────────────────────────────────────────
def _transition(slide) -> None:
    el = etree.SubElement(slide._element, f"{{{_NS}}}transition", spd="med")
    etree.SubElement(el, f"{{{_NS}}}fade")


def _animate(slide, shapes: List[Any]) -> None:
    """Fade each shape in, one after another, automatically when the slide appears."""
    if not shapes:
        return
    pars, bld, tid = [], [], 4
    for k, shp in enumerate(shapes):
        sid = shp.shape_id
        pars.append(
            f'<p:par><p:cTn id="{tid}" fill="hold"><p:stCondLst><p:cond delay="{k * 450}"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{tid + 1}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" '
            f'nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{tid + 2}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
            f'</p:cTn><p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName>'
            f'</p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{tid + 3}" dur="500"/><p:tgtEl>'
            f'<p:spTgt spid="{sid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
            f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
        tid += 4
        if getattr(shp, "has_chart", False) or getattr(shp, "has_table", False):
            bld.append(f'<p:bldGraphic spid="{sid}" grpId="0"><p:bldAsOne/></p:bldGraphic>')
        elif getattr(shp, "has_text_frame", False) and shp.has_text_frame:
            bld.append(f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>')
    xml = (f'<p:timing xmlns:p="{_NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
           '<p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
           '<p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
           '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
           + "".join(pars) +
           '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0">'
           '<p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0">'
           '<p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
           + (f'<p:bldLst>{"".join(bld)}</p:bldLst>' if bld else "") + '</p:timing>')
    slide._element.append(etree.fromstring(xml))


# ── slide furniture ─────────────────────────────────────────────────────────
class Deck:
    def __init__(self, spec: DocSpec, pal: Palette, title: str):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.spec, self.pal, self.title = spec, pal, title
        self.count = 0
        self.animated = spec.animated

    def new(self, dark: bool = False):
        slide = self.prs.slides.add_slide(self.blank)
        self.count += 1
        _bg(slide, self.pal.dark if dark else "FFFFFF")
        _transition(slide)
        return slide

    def frame(self, slide, headline: str) -> None:
        pal = self.pal
        _box(slide, 0, 0, Inches(0.18), H, pal.primary)
        _text(slide, Inches(0.6), Inches(0.38), Inches(11.6), Inches(0.9), headline, _fit(30, headline, 48),
              pal.primary, bold=True, font=pal.heading_font, anchor=MSO_ANCHOR.MIDDLE)
        _box(slide, Inches(0.68), Inches(1.28), Inches(1.1), Inches(0.07), pal.accent)
        _text(slide, Inches(0.6), Inches(6.95), Inches(9), Inches(0.35), self.title, 10, pal.muted)
        _text(slide, Inches(12.0), Inches(6.95), Inches(0.9), Inches(0.35), str(self.count), 10, pal.muted,
              align=PP_ALIGN.RIGHT)

    def notes(self, slide, text: str) -> None:
        if text:
            slide.notes_slide.notes_text_frame.text = text

    def anim(self, slide, shapes) -> None:
        if self.animated:
            _animate(slide, [s for s in shapes if s is not None])


# ── layouts ─────────────────────────────────────────────────────────────────
def _numbers_from(slide: Dict[str, Any]) -> List[Tuple[str, str]]:
    """(big number, label) pairs from DATA, else from numbers inside the bullets."""
    out = []
    for label, value, unit in slide.get("data", [])[:4]:
        num = f"{value:,.0f}" if value >= 100 or value == int(value) else f"{value:,.1f}"
        if unit.startswith("$"):
            num, unit = "$" + num, unit[1:].strip()
        out.append((num + ("%" if unit == "%" else (" " + unit if unit else "")), label))
    if out:
        return out
    value_rx = re.compile(r"\$\s?\d[\d,.]*\s*(?:trillion|billion|million|bn|[kKmMbB]\b)?|"
                          r"\d[\d,.]*\s*(?:%|x\b|trillion|billion|million|bn\b|GWh|kWh|TWh|km|miles)")
    for b in slide.get("bullets", []):
        m = value_rx.search(b)
        if m and len(out) < 4:
            label = re.sub(r"\s+", " ", b.replace(m.group(0), "")).strip(" ,.:-=")
            label = re.sub(r"^(?:label|value)\s*[=:]?\s*", "", label, flags=re.I)
            out.append((m.group(0).strip(), label[:70]))
    return out if len(out) >= 2 else []


def _bullets(d: Deck, slide, s: Dict[str, Any]) -> List[Any]:
    pal, shapes = d.pal, []
    items = s["bullets"][:5] or [s.get("notes", "")[:140]]
    n = len(items)
    gap = min(Inches(1.35), Inches(4.9) / max(n, 1))
    top = Inches(1.75) + (Inches(4.9) - gap * n) / 2                 # centre the block vertically
    size = 22 if n <= 3 else 19
    for i, b in enumerate(items):
        y = top + gap * i
        dot = _box(slide, Inches(0.75), y + Inches(0.05), Inches(0.55), Inches(0.55), pal.primary, MSO_SHAPE.OVAL)
        _text(slide, 0, 0, 0, 0, str(i + 1), 15, "FFFFFF", bold=True, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE, shape=dot)
        t = _text(slide, Inches(1.55), y, Inches(10.8), gap, b, _fit(size, b, 100), pal.text, anchor=MSO_ANCHOR.TOP)
        shapes += [dot, t]
    return shapes


def _cards(d: Deck, slide, s: Dict[str, Any]) -> List[Any]:
    pal, shapes = d.pal, []
    items = s["bullets"][:4] or ["Key point"]
    n = len(items)
    cols = 2 if n == 4 else n
    rows = 2 if n == 4 else 1
    cw = (Inches(12.0) - Inches(0.3) * (cols - 1)) / cols
    ch = Inches(3.4) if rows == 1 else Inches(2.25)
    top = Inches(1.75) + ((Inches(4.9) - ch) / 2 if rows == 1 else 0)
    parts = []
    for item in items:
        head, body = (item.split(":", 1) + [""])[:2] if ":" in item[:50] else ("", item)
        if not body.strip():
            head, body = "", head
        parts.append((head.strip(" *"), body.strip()))
    size = min(_fit(17, b, 130 if rows == 1 else 90) for _, b in parts)        # one size for every card
    for i, (head, body) in enumerate(parts):
        x = Inches(0.65) + (cw + Inches(0.3)) * (i % cols)
        y = top + (ch + Inches(0.25)) * (i // cols)
        card = _box(slide, x, y, cw, ch, pal.light, MSO_SHAPE.ROUNDED_RECTANGLE)
        _box(slide, x, y, cw, Inches(0.09), [pal.primary, pal.secondary, pal.accent, pal.primary][i])
        lines = [f"**{head}**", body] if head else [body]
        _text(slide, 0, 0, 0, 0, lines, size, pal.text, shape=card, anchor=MSO_ANCHOR.MIDDLE)
        card.text_frame.margin_left = card.text_frame.margin_right = Inches(0.25)
        shapes.append(card)
    return shapes


def _stats(d: Deck, slide, s: Dict[str, Any], illustrative: bool) -> List[Any]:
    pal, shapes = d.pal, []
    nums = _numbers_from(s)
    if not nums:
        return _bullets(d, slide, s)
    n = len(nums)
    cw = (Inches(12.0) - Inches(0.3) * (n - 1)) / n
    for i, (num, label) in enumerate(nums):
        x = Inches(0.65) + (cw + Inches(0.3)) * i
        panel = _box(slide, x, Inches(1.9), cw, Inches(2.6), pal.light, MSO_SHAPE.ROUNDED_RECTANGLE)
        _text(slide, 0, 0, 0, 0, [num], _fit(44, num, 8), pal.accent if i % 2 == 0 else pal.secondary, bold=True,
              font=pal.heading_font, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=panel)
        lab = _text(slide, x, Inches(4.6), cw, Inches(1.0), label, 15, pal.text, align=PP_ALIGN.CENTER)
        shapes += [panel, lab]
    rest = [b for b in s["bullets"] if not any(num.split()[0] in b for num, _ in nums)][:2]
    if rest:
        shapes.append(_text(slide, Inches(0.65), Inches(5.65), Inches(12), Inches(1.1), rest, 14, pal.muted))
    if illustrative:
        _text(slide, Inches(8.4), Inches(6.62), Inches(4.5), Inches(0.3), "Illustrative figures — verify before use",
              9, pal.muted, align=PP_ALIGN.RIGHT)
    return shapes


def _chart(d: Deck, slide, s: Dict[str, Any], illustrative: bool) -> List[Any]:
    pal = d.pal
    data = s.get("data") or []
    if len(data) < 2:
        return _stats(d, slide, s, illustrative)
    cd = CategoryChartData()
    cd.categories = [lab[:24] for lab, _, _ in data]
    cd.add_series(s["title"][:30], [v for _, v, _ in data])
    pie = len(data) <= 5 and all(u == "%" for _, _, u in data) and abs(sum(v for _, v, _ in data) - 100) < 5
    kind = XL_CHART_TYPE.DOUGHNUT if pie else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(kind, Inches(0.65), Inches(1.65), Inches(7.4), Inches(5.0), cd)
    ch = gf.chart
    ch.has_legend = pie
    if pie:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        for i, pt in enumerate(ch.plots[0].series[0].points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = _c([pal.primary, pal.secondary, pal.accent, pal.muted, pal.dark][i % 5])
    else:
        ser = ch.plots[0].series[0]
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = _c(pal.primary)
        ch.plots[0].gap_width = 70
        ch.plots[0].has_data_labels = True
        ch.plots[0].data_labels.font.size = Pt(11)
        ch.value_axis.has_major_gridlines = False
        ch.value_axis.visible = False
        ch.category_axis.tick_labels.font.size = Pt(11)
    side = _text(slide, Inches(8.4), Inches(1.8), Inches(4.4), Inches(4.6), s["bullets"][:4], 15, pal.text)
    if illustrative:
        _text(slide, Inches(8.4), Inches(6.62), Inches(4.5), Inches(0.3), "Illustrative figures — verify before use",
              9, pal.muted, align=PP_ALIGN.RIGHT)
    return [gf, side]


def _timeline(d: Deck, slide, s: Dict[str, Any]) -> List[Any]:
    pal, shapes = d.pal, []
    items = s["bullets"][:5] or ["Start", "Finish"]
    n = len(items)
    _box(slide, Inches(0.9), Inches(3.55), Inches(11.5), Inches(0.08), pal.secondary)
    step = Inches(11.5) / n
    for i, it in enumerate(items):
        head, body = (it.split(":", 1) + [""])[:2] if ":" in it[:30] else (f"Step {i + 1}", it)
        cx = Inches(0.9) + step * i + step / 2
        dot = _box(slide, cx - Inches(0.22), Inches(3.37), Inches(0.44), Inches(0.44), pal.accent, MSO_SHAPE.OVAL)
        up = i % 2 == 0
        lab = _text(slide, cx - step / 2 + Inches(0.05), Inches(2.3) if up else Inches(4.05), step - Inches(0.1),
                    Inches(1.0), [f"**{head.strip()}**"], 16, pal.primary, align=PP_ALIGN.CENTER,
                    anchor=MSO_ANCHOR.BOTTOM if up else MSO_ANCHOR.TOP)
        txt = _text(slide, cx - step / 2 + Inches(0.05), Inches(1.55) if up else Inches(4.75), step - Inches(0.1),
                    Inches(0.8) if up else Inches(1.6), body.strip(), _fit(12.5, body, 90), pal.text,
                    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM if up else MSO_ANCHOR.TOP)
        shapes += [dot, lab, txt]
    return shapes


def _table(d: Deck, slide, s: Dict[str, Any]) -> List[Any]:
    pal = d.pal
    rows = [[c.strip() for c in b.split("|") if c.strip()] for b in s["bullets"] if "|" in b]
    if len(rows) < 2:
        if s.get("data"):
            rows = [["Item", "Value"]] + [[lab, f"{v:,.2f}".rstrip("0").rstrip(".") + (u or "")] for lab, v, u in s["data"]]
        else:
            return _cards(d, slide, s)
    cols = max(len(r) for r in rows)
    rows = [(r + [""] * cols)[:cols] for r in rows[:8]]
    gf = slide.shapes.add_table(len(rows), cols, Inches(0.65), Inches(1.75), Inches(12.0), Inches(0.55) * len(rows))
    tbl = gf.table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = val
            para = cell.text_frame.paragraphs[0]
            para.runs[0].font.size = Pt(14 if r == 0 else 13)
            para.runs[0].font.bold = r == 0
            para.runs[0].font.color.rgb = _c("FFFFFF" if r == 0 else pal.text)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _c(pal.primary if r == 0 else (pal.light if r % 2 else "FFFFFF"))
    return [gf]


def _statement(d: Deck, slide, s: Dict[str, Any]) -> List[Any]:
    """One key message, large, with its supporting points underneath (never an invented quotation)."""
    pal = d.pal
    items = s["bullets"] or [s.get("notes", "")]
    message = items[0].strip('"“” ')
    bar = _box(slide, Inches(0.75), Inches(1.95), Inches(0.12), Inches(2.3), pal.accent)
    big = _text(slide, Inches(1.15), Inches(1.85), Inches(11.2), Inches(2.5), message, _fit(32, message, 120),
                pal.primary, bold=True, font=pal.heading_font, anchor=MSO_ANCHOR.MIDDLE)
    shapes = [bar, big]
    support = items[1:4]
    if support:
        w = Inches(11.6) / len(support)
        for i, p in enumerate(support):
            panel = _box(slide, Inches(0.75) + w * i, Inches(4.75), w - Inches(0.2), Inches(1.75), pal.light,
                         MSO_SHAPE.ROUNDED_RECTANGLE)
            _text(slide, 0, 0, 0, 0, p, _fit(15, p, 110), pal.text, shape=panel, anchor=MSO_ANCHOR.MIDDLE)
            panel.text_frame.margin_left = panel.text_frame.margin_right = Inches(0.2)
            shapes.append(panel)
    return shapes


def _two_column(d: Deck, slide, s: Dict[str, Any]) -> List[Any]:
    pal = d.pal
    items = s["bullets"][:6]
    half = (len(items) + 1) // 2
    shapes = []
    for k, part in enumerate((items[:half], items[half:])):
        x = Inches(0.65) + Inches(6.15) * k
        height = min(Inches(4.8), Inches(0.6) + Inches(1.05) * len(part))
        panel = _box(slide, x, Inches(1.75), Inches(5.85), height, pal.light, MSO_SHAPE.ROUNDED_RECTANGLE)
        _text(slide, 0, 0, 0, 0, [f"▪  {p}" for p in part], 17, pal.text, shape=panel, anchor=MSO_ANCHOR.MIDDLE)
        panel.text_frame.margin_top = panel.text_frame.margin_left = Inches(0.3)
        shapes.append(panel)
    return shapes


# ── the deck ────────────────────────────────────────────────────────────────
def render_deck(content: Dict[str, Any], spec: DocSpec, pal: Palette, path: str) -> Dict[str, Any]:
    d = Deck(spec, pal, content["title"])
    illustrative = not content.get("data_from_research")

    # Title slide
    sl = d.new(dark=True)
    _box(sl, 0, Inches(5.9), W, Inches(1.6), pal.primary)
    _box(sl, Inches(0.8), Inches(2.35), Inches(0.12), Inches(2.1), pal.accent)
    t = _text(sl, Inches(1.15), Inches(2.2), Inches(11), Inches(1.6), content["title"], _fit(44, content["title"], 40),
              "FFFFFF", bold=True, font=pal.heading_font, anchor=MSO_ANCHOR.BOTTOM)
    sub = _text(sl, Inches(1.15), Inches(3.85), Inches(11), Inches(0.8), content.get("subtitle") or spec.topic, 20,
                pal.light)
    meta = " · ".join(x for x in (spec.sender, f"{date.today():%B %d, %Y}") if x)
    m = _text(sl, Inches(1.15), Inches(6.35), Inches(11), Inches(0.6), meta, 14, "FFFFFF")
    d.anim(sl, [t, sub, m])

    # Agenda
    sl = d.new()
    d.frame(sl, "Agenda")
    titles = [s["title"] for s in content["slides"]]
    per_col = (len(titles) + 1) // 2 if len(titles) > 5 else len(titles)
    agenda_size = min(_fit(18, t, 40) for t in titles[:12]) if titles else 18
    shapes = []
    for i, ttl in enumerate(titles[:12]):
        col, row = divmod(i, per_col)
        x, y = Inches(0.75) + Inches(6.1) * col, Inches(1.75) + Inches(4.9) / max(per_col, 1) * row
        num = _text(sl, x, y, Inches(0.7), Inches(0.5), f"{i + 1:02d}", 20, pal.accent, bold=True, font=pal.heading_font)
        lab = _text(sl, x + Inches(0.75), y + Inches(0.03), Inches(5.1), Inches(0.5), ttl, agenda_size, pal.text)
        shapes += [num, lab]
    d.anim(sl, shapes)

    layouts = {"bullets": _bullets, "cards": _cards, "timeline": _timeline, "table": _table, "statement": _statement,
               "quote": _statement, "two_column": _two_column}
    for s in content["slides"]:
        sl = d.new()
        d.frame(sl, s["title"])
        kind = s.get("kind", "bullets")
        if kind == "stats":
            shapes = _stats(d, sl, s, illustrative)
        elif kind == "chart":
            shapes = _chart(d, sl, s, illustrative)
        else:
            b = s["bullets"]
            if kind == "cards" and not (2 <= len(b) <= 4 and all(":" in x[:50] and len(x.split(":", 1)[1].split()) >= 4
                                                                   for x in b)):
                kind = "bullets"
            if kind == "timeline" and sum(bool(re.match(r"^\s*(?:\d{4}|phase \d+|stage \d+|step \d+|q[1-4]\b|"
                                                        r"year \d+|month \d+)", x, re.I)) for x in b) < max(3, len(b) - 1):
                kind = "bullets"
            if kind == "two_column" and len(b) < 4:
                kind = "bullets"
            shapes = layouts.get(kind, _bullets)(d, sl, s)
        d.notes(sl, s.get("notes", ""))
        d.anim(sl, shapes)

    if content.get("sources"):
        sl = d.new()
        d.frame(sl, "Sources")
        lines = [ln.strip() for ln in content["sources"].split("\n") if ln.strip()][:16]
        box = _text(sl, Inches(0.65), Inches(1.65), Inches(12), Inches(5.1), lines, 11, pal.muted)
        d.anim(sl, [box])

    # Closing
    sl = d.new(dark=True)
    t = _text(sl, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.3), "Thank You", 54, "FFFFFF", bold=True,
              font=pal.heading_font, align=PP_ALIGN.CENTER)
    q = _text(sl, Inches(0.9), Inches(3.7), Inches(11.5), Inches(0.8), "Questions & Discussion", 22, pal.accent,
              align=PP_ALIGN.CENTER)
    who = _text(sl, Inches(0.9), Inches(5.0), Inches(11.5), Inches(0.6), spec.sender or content["title"], 16, pal.light,
                align=PP_ALIGN.CENTER)
    d.anim(sl, [t, q, who])

    d.prs.core_properties.title = content["title"]
    d.prs.core_properties.author = spec.sender or "TOM"
    d.prs.save(path)
    return {"path": path, "slides": d.count}
