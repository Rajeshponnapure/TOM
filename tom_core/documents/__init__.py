"""Professional document generation: understand → plan → write section by section → render → verify.

    result = await create("pptx", "make a 10 slide deck on EV batteries for investors", ask)

Kinds: docx (Word), pptx (animated PowerPoint), xlsx (Excel), pdf. `ask(system, user)` is any async
model call. The file is re-opened after saving and checked (FR-008) before success is reported.
"""
from __future__ import annotations

import asyncio
import os
from typing import Any, Awaitable, Callable, Dict, Optional

from tom_core.documents.design import palette_for
from tom_core.documents.spec import DocSpec, file_name_for, understand
from tom_core.documents import writer

Ask = Callable[[str, str], Awaitable[str]]
Research = Callable[[str], Awaitable[str]]

__all__ = ["create", "understand", "DocSpec", "verify"]


def _output_path(name: str) -> str:
    from tools.document_creator import _output_path as gated          # R3: workspace write gate
    base, ext = os.path.splitext(name)
    path = gated(name)
    n = 2
    while os.path.exists(path):                                       # never overwrite an earlier file
        path = gated(f"{base}_{n}{ext}")
        n += 1
    return path


def verify(kind: str, path: str) -> Dict[str, Any]:
    """Re-open the saved file with its own library and report what is really inside."""
    if not os.path.isfile(path) or os.path.getsize(path) < 1000:
        return {"ok": False, "detail": "file missing or nearly empty"}
    try:
        if kind == "docx":
            from docx import Document
            d = Document(path)
            n = sum(len(p.text.split()) for p in d.paragraphs) + sum(
                len(c.text.split()) for t in d.tables for row in t.rows for c in row.cells)
            return {"ok": n > 80, "words": n, "detail": f"{n:,} words"}
        if kind == "pptx":
            from pptx import Presentation
            p = Presentation(path)
            anim = sum(1 for s in p.slides if s._element.find(
                "{http://schemas.openxmlformats.org/presentationml/2006/main}timing") is not None)
            return {"ok": len(p.slides) >= 3, "slides": len(p.slides), "animated_slides": anim,
                    "detail": f"{len(p.slides)} slides, {anim} animated"}
        if kind == "xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(path)
            cells = sum(ws.max_row * ws.max_column for ws in wb.worksheets)
            return {"ok": cells > 10, "sheets": wb.sheetnames, "detail": f"{len(wb.sheetnames)} sheets"}
        if kind == "pdf":
            with open(path, "rb") as fh:
                data = fh.read()
            pages = data.count(b"/Type /Page") - data.count(b"/Type /Pages")
            return {"ok": data.startswith(b"%PDF") and pages >= 1, "pages": pages, "detail": f"{pages} pages"}
    except Exception as exc:
        return {"ok": False, "detail": f"could not be re-opened: {type(exc).__name__}: {exc}"}
    return {"ok": False, "detail": "unknown kind"}


async def create(kind: str, request: str, ask: Ask, *, parsed: Optional[Dict[str, Any]] = None,
                 research: Optional[Research] = None,
                 progress: Optional[Callable[[str], None]] = None) -> Dict[str, Any]:
    spec = understand(kind, request, parsed)
    notes, sources = "", ""
    if research is not None and spec.research:
        try:
            got = await research(spec.topic)
            notes, sources = got if isinstance(got, tuple) else (got, "")
        except Exception:
            notes, sources = "", ""
    pal = palette_for(spec.theme, spec.topic)
    if kind == "pptx":
        content = await writer.write_deck(spec, ask, notes, progress)
    elif kind == "xlsx":
        content = await writer.write_workbook(spec, ask, notes)
    else:
        content = await writer.write_document(spec, ask, notes, progress)

    if sources and kind in ("docx", "pdf") and not content.get("letter"):
        content["sections"].append({"heading": "Sources", "issues": [], "blocks": [
            {"type": "paragraph", "text": "Facts in this document were drawn from these web sources (retrieved "
                                          f"{__import__('datetime').date.today():%B %d, %Y}):"},
            {"type": "bullets", "items": [ln.strip() for ln in sources.split("\n") if ln.strip()]}]})
    if sources and kind == "pptx":
        content["sources"] = sources
    path = _output_path(file_name_for(spec, content.get("subject") or content.get("title", "")))
    if kind == "pptx":
        from tom_core.documents.pptx_render import render_deck as render
    elif kind == "xlsx":
        from tom_core.documents.xlsx_render import render_workbook as render
    elif kind == "pdf":
        from tom_core.documents.pdf_render import render_pdf as render
    else:
        from tom_core.documents.docx_render import render_document as render
    await asyncio.to_thread(render, content, spec, pal, path)
    check = verify(kind, path)
    label = {"docx": "Word document", "pptx": "PowerPoint deck", "xlsx": "Excel workbook", "pdf": "PDF"}[kind]
    what = spec.genre if spec.genre not in ("presentation", "workbook") else ""
    issues = content.get("issues", [])
    phrase = f"{what} {label}" if what else label
    article = "an" if phrase[:1].lower() in "aeiou" else "a"
    msg = (f"Created {article} {phrase}: \"{content['title']}\" ({check.get('detail', '')}"
           f"{', ' + pal.name + ' theme' if kind == 'pptx' else ''}).\nSaved to: {path}")
    if kind == "pptx" and spec.animated:
        msg += "\nSlides use fade transitions and elements animate in one after another; speaker notes included."
    if notes:
        msg += "\nFacts were drawn from a live web search."
    elif kind in ("pptx", "xlsx"):
        msg += "\nNumbers not taken from your request or a web search are marked as illustrative."
    if issues:
        msg += "\nQuality notes: " + "; ".join(issues[:4])
    return {"status": "success" if check["ok"] else "error", "path": path, "paths": [path],
            "message": msg if check["ok"] else f"The {label} was saved but failed its check: {check['detail']}. {path}",
            "verified": check["ok"], "evidence": check.get("detail", ""), "response_type": f"document_{kind}",
            "quality_issues": issues, "spec": {"genre": spec.genre, "topic": spec.topic, "theme": pal.name}}
