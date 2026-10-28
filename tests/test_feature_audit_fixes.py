"""Regressions for the defects the live feature audit found on 2026-10-27 (every quick action run through
the real agent and model). Fakes only — no network, Blender, Chrome or model calls.
"""
import asyncio
import os
import subprocess
from datetime import datetime, timedelta, timezone
from email.utils import format_datetime

import pytest

from tom_core.documents.design import palette_for
from tom_core.documents.spec import understand


# ── documents: exact deck length, one-page documents, Excel formulas ───────

@pytest.mark.parametrize("request_text,content,total", [
    ("5 slides introducing our bakery app to investors", 3, 5),
    ("a 10 slide deck on EV batteries with latest research", 6, 10),
    ("a 12 slide deck about coffee", 8, 12),
])
def test_a_requested_slide_count_is_the_whole_deck(request_text, content, total):
    spec = understand("pptx", request_text)
    assert spec.slides == content and spec.extras["slides_total"] == total


def test_the_rendered_deck_has_exactly_the_requested_slides(tmp_path):
    from pptx import Presentation
    from tom_core.documents.pptx_render import render_deck
    spec = understand("pptx", "5 slides introducing our bakery app to investors")
    content = {"title": "Bakery App", "subtitle": "Investor intro", "sources": "[1] Source one\n[2] Source two",
               "slides": [{"title": f"Point {i}", "kind": "bullets", "bullets": ["One idea here", "Another idea"],
                           "notes": ""} for i in range(spec.slides)]}
    path = str(tmp_path / "deck.pptx")
    render_deck(content, spec, palette_for("", spec.topic), path)
    prs = Presentation(path)
    assert len(prs.slides) == 5
    assert "Source one" in prs.slides[-1].notes_slide.notes_text_frame.text      # sources kept, in the notes


@pytest.mark.parametrize("text,pages", [("a one-page project proposal", 1), ("a two page summary of AI", 2),
                                        ("3 pages report on X", 3), ("a report on X", 0)])
def test_page_counts_in_words_are_understood(text, pages):
    assert understand("docx", text).pages == pages


def test_a_one_page_document_has_no_cover_or_contents(tmp_path):
    from docx import Document
    from tom_core.documents.docx_render import render_document
    spec = understand("docx", "a one-page project proposal for a bakery app")
    content = {"title": "Bakery App Proposal", "subtitle": "", "sections": [
        {"heading": "Summary", "blocks": [{"type": "paragraph", "text": "We propose an ordering app."}]},
        {"heading": "Plan", "blocks": [{"type": "bullets", "items": ["Build: four weeks", "Launch: week five"]}]}]}
    path = str(tmp_path / "p.docx")
    render_document(content, spec, palette_for(), path)
    texts = [p.text for p in Document(path).paragraphs]
    assert "Contents" not in texts and "Bakery App Proposal" in texts and "Summary" in texts
    assert sum(1 for t in texts if not t.strip()) < 3                    # no blank lines padding a cover page


def test_amount_columns_get_live_totals_even_when_labelled_text(tmp_path):
    from openpyxl import load_workbook
    from tom_core.documents.xlsx_render import currency_format, infer_kinds, render_workbook
    rows = [["Salary", "₹50,000", "Monthly"], ["Freelance", "Rs 12,000", "Monthly"], ["Rent", "₹ 18,500", "Monthly"]]
    cols = [{"name": "Category", "type": "text"}, {"name": "Amount", "type": "text"}, {"name": "Frequency", "type": "text"}]
    assert infer_kinds(cols, rows) == ["text", "currency", "text"]
    assert currency_format(rows) == '"₹"#,##0.00' and currency_format([["a", "12"]]) == "#,##0.00"
    spec = understand("xlsx", "a monthly household budget")
    path = str(tmp_path / "b.xlsx")
    render_workbook({"title": "Budget", "sheets": [{"name": "Income", "columns": cols, "rows": rows}]},
                    spec, palette_for(), path)
    wb = load_workbook(path)
    ws = wb["Income"]
    formulas = [c.value for r in ws.iter_rows() for c in r if isinstance(c.value, str) and c.value.startswith("=")]
    assert "=SUM(B5:B7)" in formulas and "Summary" in wb.sheetnames
    assert ws["B5"].value == 50000 and ws["B5"].number_format == '"₹"#,##0.00'


# ── Blender ────────────────────────────────────────────────────────────

def test_a_house_with_lights_is_a_house_with_lights_not_a_lighting_demo():
    from tools.blender_control import BlenderControl
    b = BlenderControl.__new__(BlenderControl)
    script = b.generate_script("a low-poly house with lights")
    assert "TOM_Roof" in script and "TOM_PorchLight" in script and "TOM_Window_" in script
    assert "TOM_PorchLight" not in b.generate_script("a house") and "TOM_Roof" in b.generate_script("a cottage")


def test_blender_runs_without_the_owners_addons_and_keeps_a_saved_scene(monkeypatch, tmp_path):
    from tools import blender_control
    from tools.blender_control import BlenderControl
    b = BlenderControl.__new__(BlenderControl)
    b.blender_path, b._has_bpy = "blender.exe", False
    calls = []

    def fake_run(cmd, **kw):
        calls.append(cmd)
        script = open(cmd[-1], encoding="utf-8").read()
        blend = script.split("save_as_mainfile(filepath=")[1].split(")")[0].strip("'\"")
        open(blend, "w").close()                                  # the scene was saved…
        return subprocess.CompletedProcess(cmd, 1, "", "addon unregister failed")   # …then exit 1 at shutdown
    monkeypatch.setattr(blender_control.subprocess, "run", fake_run)
    monkeypatch.setattr(blender_control, "project_path_str", lambda *p: str(tmp_path.joinpath(*p)))
    res = b.generate_and_execute("a house")
    assert "--factory-startup" in calls[0]
    assert res["status"] == "success" and res["path"].endswith("scene.blend")
    stage = open(calls[0][-1], encoding="utf-8").read()
    assert "camera_add" in stage and "TRACK_TO" in stage                    # every scene is framed and lit


def test_a_failing_engine_is_reported_not_replaced_by_a_skill_answer():
    from agent import TomAgent
    from tom_core import skill_modes

    class Router:
        def detect(self, c):
            return None

        async def execute(self, c, key=None):
            return {"status": "error", "message": "Blender: exited with code 1"}
    a = TomAgent.__new__(TomAgent)
    a.engine_router = Router()
    res = asyncio.run(a._run_skill_mode(skill_modes.MODES["blender"], "a low-poly house"))
    assert res["status"] == "error" and "Blender" in res["message"]


# ── research is grounded; analysis is structured ───────────────────────

def test_specialist_research_uses_web_results_and_lists_sources(monkeypatch):
    from tools import agent_orchestrator as ao, web_search
    hits = [web_search.SearchHit(title="HubSpot CRM free plan", url="https://example.org/a", source="Ex", snippet="x")]
    monkeypatch.setattr(web_search, "search", lambda q, n=8: hits)
    seen = {}

    async def ask(system, user):
        seen["system"], seen["user"] = system, user
        return "HubSpot has a free plan [1]."
    orch = ao.AgentOrchestrator.__new__(ao.AgentOrchestrator)
    orch._agents = {r: type("A", (), {"name": r.value})() for r in ao.AgentRole}
    orch._ask, orch.llm, orch.tom_agent = ask, object(), None
    res = asyncio.run(orch._specialist(ao.AgentRole.CMO, "compare the top 3 CRM tools for small businesses"))
    assert "ONLY on the numbered" in seen["system"] and "HubSpot CRM free plan" in seen["user"]
    assert "Sources:" in res["message"] and "https://example.org/a" in res["message"]


def test_without_web_results_research_must_not_state_figures(monkeypatch):
    from tools import agent_orchestrator as ao, web_search
    monkeypatch.setattr(web_search, "search", lambda q, n=8: [])
    seen = {}

    async def ask(system, user):
        seen["system"] = system
        return "Compare pricing tiers; figures need checking."
    orch = ao.AgentOrchestrator.__new__(ao.AgentOrchestrator)
    orch._agents = {r: type("A", (), {"name": r.value})() for r in ao.AgentRole}
    orch._ask, orch.llm, orch.tom_agent = ask, object(), None
    asyncio.run(orch._specialist(ao.AgentRole.CMO, "compare CRM pricing"))
    assert "do not state specific prices" in seen["system"]


def test_analyze_mode_gives_a_structured_analysis():
    from agent import TomAgent
    from tom_core import skill_modes
    seen = {}

    class Resp:
        content = "## Summary\n" + "word " * 120

    async def invoke(prompt, payload, label, llm=None):
        seen.update(payload)
        return Resp()
    a = TomAgent.__new__(TomAgent)
    a._invoke_llm, a.llm, a.engine_router = invoke, None, None
    res = asyncio.run(a._run_skill_mode(skill_modes.MODES["analyze"], "pros and cons of remote work for a startup"))
    assert res["response_type"] == "analysis" and res["status"] == "success"
    assert "## Recommendation" in seen["rules"] and "## Risks" in seen["rules"]


# ── news, diagnostics, dependencies, updates, ML, website safety ───────

def test_news_briefings_are_fresh_linked_and_deduplicated():
    from tools import news_briefing
    from tools.web_search import SearchHit
    now = datetime(2026, 10, 27, 12, tzinfo=timezone.utc)

    def fetch(query, limit):
        return [SearchHit("Big AI launch", "https://n/1", source="A", date=format_datetime(now - timedelta(hours=2))),
                SearchHit("Big AI launch", "https://n/1b", source="B", date=format_datetime(now - timedelta(hours=3))),
                SearchHit("Top 10 AI tools in 2022", "https://n/old", source="C",
                          date=format_datetime(now - timedelta(days=900))),
                SearchHit("No link", "", source="D", date=format_datetime(now))]
    brief = news_briefing.build("tech", fetch=fetch, now=now)
    text = brief["message"]
    assert brief["status"] == "success" and text.count("Big AI launch") == 1
    assert "2022" not in text and "https://n/1" in text and "No link" not in text
    assert news_briefing.build("tech", fetch=lambda q, n: [], now=now)["status"] == "error"


@pytest.mark.parametrize("text,what", [("why is my RAM usage high?", "ram"), ("what is using my cpu", "cpu"),
                                       ("which app is eating my memory", "ram"), ("my pc is so slow", "all"),
                                       ("why is my disk full", "disk"), ("what is RAM", None),
                                       ("how much ram do i need for gaming", None)])
def test_diagnostic_questions_are_recognised(text, what):
    from tom_core import diagnostics
    from tools.command_router import CommandRouter
    assert diagnostics.wants(text) == what
    assert (CommandRouter().route(text).handler == "execute_diagnostics") is (what is not None)


def test_the_diagnostics_report_is_measured_and_closes_nothing():
    from tom_core import diagnostics
    res = diagnostics.report("ram")
    assert res["status"] == "success" and "RAM" in res["message"] and "Using the most memory" in res["message"]
    assert "haven't closed anything" in res["message"]


def test_a_clean_environment_is_not_reported_as_conflicting(monkeypatch):
    from tools.dependency_manager import DependencyManager
    dm = DependencyManager()
    monkeypatch.setattr(dm, "_package_is_installed", lambda *a, **k: False)
    monkeypatch.setattr(dm, "_run_cmd", lambda cmd, **k: {"success": True, "stdout": "No broken requirements found."})
    res = dm.check_conflicts()
    assert res["result"]["has_conflicts"] is False and "No broken requirements" in res["message"]
    monkeypatch.setattr(dm, "_run_cmd", lambda cmd, **k: {
        "success": False, "stdout": "torch 2.1 has requirement numpy<2, but you have numpy 2.1."})
    res = dm.check_conflicts()
    assert res["result"]["has_conflicts"] is True and "numpy" in res["message"]


def test_update_check_compares_with_the_remote(tmp_path):
    from tools.auto_update import AutoUpdate

    def git(*args, cwd):
        subprocess.run(["git", "-c", "user.email=t@t", "-c", "user.name=t", *args], cwd=cwd, check=True,
                       capture_output=True)
    remote, local = tmp_path / "remote", tmp_path / "local"
    remote.mkdir()
    git("init", "-q", "-b", "main", cwd=remote)
    (remote / "a.txt").write_text("1")
    git("add", ".", cwd=remote)
    git("commit", "-qm", "first", cwd=remote)
    git("clone", "-q", str(remote), str(local), cwd=tmp_path)
    (remote / "a.txt").write_text("2")
    git("commit", "-qam", "Add a new feature", cwd=remote)
    au = AutoUpdate.__new__(AutoUpdate)
    au.project_dir = str(local)
    res = au.check_git_updates()
    assert res["behind"] == 1 and res["ahead"] == 0 and "Add a new feature" in res["message"]
    git("pull", "-q", cwd=local)
    assert au.check_git_updates()["status"] == "up_to_date"


def test_ml_predicts_what_was_asked():
    from tools.engine_router import EngineRouter
    msg = EngineRouter()._run_ml("train a regression on X=[[1],[2],[3],[4]] y=[2,4,6,8] and predict for 5")["message"]
    assert "Predictions:" in msg and "x = 5.0 → 10" in msg and "Fitted line: y = 2·x" in msg


def test_a_page_budget_is_enforced():
    from tom_core.documents.writer import fit_length
    long = " ".join(["This sentence has exactly eight words in it."] * 60)            # 480 words
    sections = [{"heading": f"S{i}", "blocks": [{"type": "paragraph", "text": long},
                                                {"type": "bullets", "items": [long] * 8},
                                                {"type": "table", "headers": ["a"], "rows": [["x"]] * 20}]}
                for i in range(3)]
    fit_length(sections, 500)
    total = sum(len(str(b.get("text", "")).split()) + sum(len(i.split()) for i in b.get("items", []))
                for s in sections for b in s["blocks"])
    assert total <= 560 and all(s["blocks"][0]["text"].endswith(".") for s in sections)
    assert understand("docx", "a one-page proposal").target_sections == 3


def test_double_escaped_text_from_a_plan_gets_real_line_breaks():
    from tom_core.planning.planner import _clean_args
    args = _clean_args({"path": "poems/poem1.txt", "content": "The sun sets slow,\\nthe sky turns gold."})
    assert args["content"] == "The sun sets slow,\nthe sky turns gold."
    assert _clean_args({"content": "real\nlines and a literal \\n"})["content"] == "real\nlines and a literal \\n"


def test_a_code_fix_is_a_minimal_change():
    from tom_core.coding.workflow import keep_formatting
    original = "def add(a, b):\n    return a - b\n\n\ndef mul(a, b):\n    return a * b\n"
    proposed = "def add(a, b):\n    return a + b\ndef mul(a, b):\n    return a * b"           # model's rewrite
    assert keep_formatting(original, proposed) == original.replace("a - b", "a + b")
    assert keep_formatting("if x:\nprint(1)\n", "if x:\n    print(1)\n") == "if x:\n    print(1)\n"
    assert keep_formatting("a = 1\r\nb = 2\r\n", "a = 1\nb = 3\n") == "a = 1\r\nb = 3\r\n"


def test_website_safety_tests_https_instead_of_reading_the_typed_address(monkeypatch):
    from safety.guards import SafetyGuards
    from tools import engine_router
    monkeypatch.setattr(engine_router, "_probe_https", lambda d: "valid HTTPS certificate; responded HTTP 200")
    er = engine_router.EngineRouter()
    er._agent = type("A", (), {"safety": SafetyGuards()})()
    msg = er._run_websafety("is example.com safe to open?")["message"]
    assert "no https" not in msg and "valid https certificate" in msg
    monkeypatch.setattr(engine_router, "_probe_https", lambda d: "certificate problem — could not be verified")
    msg = er._run_websafety("is secure-login-verify1234.com safe?")["message"]
    assert "certificate error" in msg and "score 0/100" in msg


def test_plans_name_known_folders_by_short_tokens_expanded_exactly():
    """qwen3:8b copied the long workspace path wrongly and the folder landed outside the workspace."""
    from tom_core.planning.planner import _clean_args, _known_folders
    ws = _known_folders()["workspace"]
    assert _clean_args({"path": "<workspace>/poems/poem1.txt"})["path"] == f"{ws}/poems/poem1.txt"
    assert _clean_args({"path": "<Desktop>" + chr(10) + "otes.txt"})["path"] == _known_folders()["desktop"] + "/notes.txt"
    assert _clean_args({"path": "C:/Users/x/file.txt"})["path"] == "C:/Users/x/file.txt"


def test_a_named_folder_missing_from_the_plan_is_put_back():
    from tom_core.planning.planner import Step, repair_named_folder
    ws = "C:/W"
    steps = [Step(id=1, tool="fs.make_dir", args={"path": ws}),
             Step(id=2, tool="fs.write_text", args={"path": f"{ws}/poem1.txt", "content": "x"}),
             Step(id=3, tool="fs.write_text", args={"path": "C:/elsewhere/keep.txt", "content": "y"})]
    repair_named_folder("create a folder called poems with 3 short poems in my workspace", steps)
    assert steps[0].args["path"] == "C:/W/poems" and steps[1].args["path"] == "C:/W/poems/poem1.txt"
    assert steps[2].args["path"] == "C:/elsewhere/keep.txt"
    ok = [Step(id=1, tool="fs.make_dir", args={"path": "C:/W/poems"})]
    repair_named_folder("create a folder called poems", ok)
    assert ok[0].args["path"] == "C:/W/poems"                                  # already right: untouched
