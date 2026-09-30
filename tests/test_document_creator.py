"""create_presentation() must survive whatever shape the LLM actually returns for slide content.

Real bug, found by physically generating a deck: the prompt in agent.py asks the LLM for
`"content": ["Bullet 1", "Bullet 2"]`, but nothing validated the response before it reached
_add_bullet_box(), which does `for bullet in bullets[:8]`. When a slide's content came back as
a plain string (an ordinary LLM deviation -- a paragraph instead of a bullet list) that iterated
per CHARACTER, rendering one bullet per letter instead of one bullet for the sentence.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest

pytest.importorskip("pptx")

from tools.document_creator import create_presentation


def _slide_texts(fpath):
    from pptx import Presentation
    prs = Presentation(fpath)
    out = []
    for slide in prs.slides:
        out.append([shape.text_frame.text for shape in slide.shapes if shape.has_text_frame])
    return out


def test_a_string_content_field_becomes_one_bullet_not_one_per_letter(tmp_path):
    fpath = str(tmp_path / "deck.pptx")
    result = create_presentation(
        title="Test Deck",
        slides_content=[{"title": "Intro", "content": "Welcome"}],
        filename=fpath,
    )
    assert result["status"] == "success", result
    texts = _slide_texts(fpath)
    # slide 0 is the title slide; slide 1 is the one content slide
    flat = " ".join(texts[1])
    assert "Welcome" in flat
    assert "W e l c o m e" not in flat and "W\ne\nl" not in flat


def test_a_real_bullet_list_still_renders_as_separate_bullets(tmp_path):
    fpath = str(tmp_path / "deck2.pptx")
    result = create_presentation(
        title="Test Deck",
        slides_content=[{"title": "Details", "content": ["Point A", "Point B", "Point C"]}],
        filename=fpath,
    )
    assert result["status"] == "success", result
    flat = " ".join(_slide_texts(fpath)[1])
    assert "Point A" in flat and "Point B" in flat and "Point C" in flat


def test_a_non_dict_slide_entry_does_not_crash_the_whole_deck(tmp_path):
    fpath = str(tmp_path / "deck3.pptx")
    result = create_presentation(
        title="Test Deck",
        slides_content=[{"title": "Normal", "content": ["ok"]}, "a bare string slide"],
        filename=fpath,
    )
    assert result["status"] == "success", result
    assert os.path.isfile(fpath)
