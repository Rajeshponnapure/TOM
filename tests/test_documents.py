"""Professional document pipeline (tom_core.documents) — understanding, writing, rendering, verification.

Owner report 2026-10-22: generated Word/PPT/Excel/PDF files looked like random templates. These tests
pin the structure that fixed it; a scripted model stands in for the LLM.
"""
import asyncio
import json
import os
import re
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tom_core import documents  # noqa: E402
from tom_core.documents import spec as S, writer as W  # noqa: E402

PARA = ("Hospitals use machine learning to triage scans faster, and clinicians review every flagged case. "
        "Adoption is growing where data quality is high, while smaller clinics still face cost barriers. ") * 3


def fake_model(bad_first=False):
    calls = {"n": 0}

    async def ask(system, user):
        calls["n"] += 1
        if "plan professional documents" in system:
            return json.dumps({"title": "AI in Healthcare: Opportunities and Risks", "subtitle": "A review",
                               "sections": [{"heading": h, "purpose": "x"} for h in
                                            ("Executive Summary", "Current Landscape", "Key Findings", "Recommendations")]})
        if "plan professional slide decks" in system:
            return json.dumps({"title": "EV Batteries", "slides": [
                {"title": "Market Size Today", "kind": "stats"}, {"title": "Leading Chemistries", "kind": "cards"},
                {"title": "Adoption Roadmap", "kind": "timeline"}, {"title": "Key Risks", "kind": "bullets"}]})
        if "slide content" in system:
            heads = re.findall(r"^- (.+?) \(kind: (\w+)\)", user, re.M)
            out = []
            for title, kind in heads:
                if kind == "stats":
                    pts = ["Global battery demand reached 750 GWh in 2023", "Prices fell 14% year over year"]
                elif kind == "cards":
                    pts = ["LFP: Lower cost chemistry favored for standard range cars and buses today",
                           "NMC: Higher energy density chemistry used in long range premium vehicles",
                           "Solid state: Emerging chemistry promising faster charging and better safety"]
                elif kind == "timeline":
                    pts = ["2024: Gigafactories ramp in Europe", "2026: Sodium-ion enters entry models",
                           "2028: Solid-state pilot lines scale", "2030: Recycling covers a fifth of supply"]
                else:
                    pts = ["Raw material prices remain volatile across lithium and nickel markets",
                           "Policy shifts can change demand in key regions quickly",
                           "Recycling capacity lags the growth of retired battery packs"]
                out.append(f"=== SLIDE: {title}\n" + "\n".join(f"- {p}" for p in pts) +
                           "\nNOTES: Explain the point with one example.")
            return "\n".join(out)
        if "Excel workbooks" in system:
            return json.dumps({"title": "Household Budget", "sheets": [
                {"name": "Expenses", "description": "Monthly spend", "columns": [
                    {"name": "Category", "type": "text"}, {"name": "Amount", "type": "currency"}],
                 "rows": [["Rent", 1500], ["Groceries", 420.5], ["Transport", "180"], ["Total", 2100.5]],
                 "chart": {"type": "bar", "label_column": "Category", "value_columns": ["Amount"]}}]})
        if "SUBJECT:" in system:
            return ("SUBJECT: Request for Sick Leave\nBODY:\nOctober 22, 2026\nDear Manager,\n"
                    "I am writing to request sick leave for tomorrow, as I have developed a fever and my doctor has "
                    "advised rest. I will stay reachable by email for anything urgent.\n\n"
                    "I have briefed a colleague on my open tasks so that work continues without delay, and I expect "
                    "to return the following day.\n\nSincerely,\nRajesh")
        if bad_first and calls["n"] == 2:
            return "Here is the section: [Insert key findings here]"
        return PARA + "\n\n| Area | Effect |\n|---|---|\n| Imaging | Faster reads |\n| Triage | Shorter waits |"
    ask.calls = calls
    return ask


@pytest.fixture(autouse=True)
def workspace(tmp_path, monkeypatch):
    monkeypatch.setenv("TOM_WORKSPACE_DIR", str(tmp_path))
    from tools import document_creator
    from pathlib import Path
    monkeypatch.setattr(document_creator, "OUTPUT_DIR", Path(tmp_path))
    return tmp_path


def run(coro):
    return asyncio.run(coro)


# ── understanding ───────────────────────────────────────────────────────────
def test_understanding_reads_genre_audience_counts_and_recipient():
    d = S.understand("pptx", "Create a professional animated 12 slide presentation on EV batteries for investors in a dark theme")
    # "12 slide" = a 12-slide deck: title, agenda, sources, closing + 8 content slides (2026-10-27)
    assert (d.genre, d.audience, d.extras["slides_total"], d.theme, d.animated) == \
        ("presentation", "investors", 12, "dark", True)
    assert d.slides == 8
    assert "ev batteries" in d.topic.lower()
    letter = S.understand("docx", "write a sick leave letter to my manager for tomorrow")
    assert letter.genre == "letter" and letter.recipient == "manager"
    assert S.understand("docx", "write a 3 page report on solar energy").pages == 3
    assert not S.understand("pptx", "make a deck on cats without animations").animated


def test_parser_handles_tables_bullets_code_and_links():
    blocks = W.parse_blocks("Intro text.\n\n- one\n- two\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n"
                            "```python\nx = 1\n```\nSee [docs](https://example.org).")
    assert [b["type"] for b in blocks] == ["paragraph", "bullets", "table", "code", "paragraph"]
    assert blocks[2]["rows"] == [["1", "2"]] and "https://example.org" in blocks[-1]["text"]


def test_placeholders_are_detected():
    assert W.problems("Text [Insert company name] here " * 20, 10)
    assert not W.problems(PARA, 20)


# ── end to end, one per format ──────────────────────────────────────────────
def test_word_report_has_cover_contents_sections_and_no_placeholders():
    from docx import Document
    ask = fake_model(bad_first=True)
    res = run(documents.create("docx", "Write a detailed professional report on AI in healthcare", ask))
    assert res["status"] == "success" and res["verified"]
    doc = Document(res["path"])
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Contents" in text and "1. Executive Summary" in text and "4. Recommendations" in text
    assert "[Insert" not in text and len(doc.tables) >= 4                       # the placeholder draft was rewritten
    assert doc.sections[-1].footer.paragraphs[0].text.startswith("Page")


def test_letter_has_business_layout_and_no_duplicate_date_or_signoff():
    from docx import Document
    res = run(documents.create("docx", "write a sick leave letter to my manager for tomorrow because of fever", fake_model()))
    paras = [p.text for p in Document(res["path"]).paragraphs if p.text.strip()]
    assert paras[0] == S.owner_name() or paras[0]
    assert "Subject: Request for Sick Leave" in paras and "Dear Manager," in paras
    assert sum(p.startswith("Sincerely") for p in paras) == 1
    assert sum(bool(re.match(r"October \d+, 2026", p)) for p in paras) <= 1
    assert os.path.basename(res["path"]) == "request_sick_leave_letter.docx"


def test_deck_is_animated_with_designed_layouts_notes_and_honest_numbers():
    from pptx import Presentation
    res = run(documents.create("pptx", "Make a professional 6 slide presentation on EV batteries for investors", fake_model()))
    assert res["status"] == "success"
    prs = Presentation(res["path"])
    assert len(prs.slides) == 6                      # exactly as asked: title, 4 content, closing (no agenda < 8)
    ns = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    assert all(s._element.find(f"{ns}transition") is not None for s in prs.slides)
    assert all(s._element.find(f"{ns}timing") is not None for s in prs.slides)
    texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
    assert "Illustrative figures" in texts[1] and "750 GWh" in texts[1]          # stats slide
    assert "2026" in texts[3] and "Sodium-ion" in texts[3]                       # timeline slide
    assert prs.slides[2].notes_slide.notes_text_frame.text                       # speaker notes


def test_workbook_has_formulas_formats_chart_and_no_model_total_row():
    from openpyxl import load_workbook
    res = run(documents.create("xlsx", "Create an Excel sheet to track my monthly household budget", fake_model()))
    wb = load_workbook(res["path"])
    assert wb.sheetnames == ["Summary", "Expenses"]
    ws = wb["Expenses"]
    col_a = [c.value for c in ws["A"]]
    assert col_a.count("Total") == 1                                             # model's own total row dropped
    formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert "=SUM(B5:B7)" in formulas and ws["B6"].value == 420.5 and ws["B7"].value == 180
    # no currency symbol in the figures → no invented "$" (₹/€/£/$ are used when the data shows them)
    assert ws["B5"].number_format == "#,##0.00" and ws._charts and ws.freeze_panes == "A5"
    assert "illustration" in " ".join(str(c.value) for c in ws["A"] if c.value)


def test_pdf_guide_with_sources_section_when_researched():
    async def research(topic):
        return ("[1] Python 3.13 released (python.org)", "[1] Python 3.13 released — python.org\n    https://www.python.org")
    res = run(documents.create("pdf", "Write a report on learning Python for beginners", fake_model(), research=research))
    assert res["status"] == "success" and res["verified"]
    data = open(res["path"], "rb").read()
    assert data.startswith(b"%PDF") and b"Sources" in data or res["evidence"].endswith("pages")
    assert "web search" in res["message"]
