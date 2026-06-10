import os
import subprocess
from typing import Optional, Dict, Any, List

from safety.guards import SafetyGuards
from tools.project_paths import PROJECT_ROOT


class FileTools:
    def __init__(self):
        self.project_root = str(PROJECT_ROOT)
        self.config_dir = os.path.join(self.project_root, "config")
        os.makedirs(self.config_dir, exist_ok=True)

        # Default output directory for generated documents
        self.output_dir = os.path.join(self.project_root, "output")
        os.makedirs(self.output_dir, exist_ok=True)

    # ── File Writing ──────────────────────────────────────────────────────

    async def write_file(self, filepath: str, content: str) -> Dict[str, Any]:
        safety = SafetyGuards()
        if not safety.check_file_path_safe(filepath):
            return {"status": "error", "message": "Access denied: unsafe file path"}
        try:
            dir_part = os.path.dirname(filepath)
            if dir_part:
                os.makedirs(dir_part, exist_ok=True)
            with open(filepath, "w", encoding="utf-8") as f:
                f.write(content)
            return {"status": "success", "message": f"Wrote to file: {filepath}", "bytes_written": len(content)}
        except Exception as exc:
            return {"status": "error", "message": f"Failed to write file: {exc}"}

    async def read_file(self, filepath: str) -> Dict[str, Any]:
        try:
            with open(filepath, "r", encoding="utf-8") as f:
                return {"status": "success", "content": f.read()}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    async def list_directory(self, path: str) -> Dict[str, Any]:
        try:
            return {"status": "success", "files": os.listdir(path)}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    def open_file(self, filepath: str):
        try:
            os.startfile(filepath)
        except Exception:
            pass

    # ── Website Creation ──────────────────────────────────────────────────

    async def create_website(self, project_name: str, index_content: Optional[str] = None,
                             css_content: Optional[str] = None) -> Dict[str, Any]:
        try:
            target_dir = os.path.join("output", project_name)
            os.makedirs(target_dir, exist_ok=True)
            final_html = index_content or "<!DOCTYPE html>\n<html>\n<head><title>Page</title></head>\n<body>\n<h1>Welcome</h1>\n</body>\n</html>"
            final_css = css_content or "body { margin: 0; font-family: Arial, sans-serif; }"
            index_path = os.path.join(target_dir, "index.html")
            css_path = os.path.join(target_dir, "style.css")
            with open(index_path, "w", encoding="utf-8") as f:
                f.write(final_html)
            with open(css_path, "w", encoding="utf-8") as f:
                f.write(final_css)
            return {"status": "success", "message": f"Website created in {target_dir}/",
                    "files_created": ["index.html", "style.css"], "path": target_dir}
        except Exception as e:
            return {"status": "error", "message": f"Failed to create website: {str(e)}"}

    # ── PROFESSIONAL WORD DOCUMENT (.docx) ───────────────────────────────

    async def create_word_document(self, file_name: str, title: str = "Document",
                                   author: str = "TOM AI", content_sections: Optional[List[Dict]] = None,
                                   template: str = "professional") -> Dict[str, Any]:
        """
        Create a professional Word document.
        content_sections: list of dicts like:
          {"type": "heading", "text": "Introduction", "level": 1},
          {"type": "paragraph", "text": "This is body text..."},
          {"type": "bullet", "text": "Item 1"},
          {"type": "table", "headers": ["Col1","Col2"], "rows": [["a","b"]]}
        """
        try:
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor, Cm
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.enum.table import WD_TABLE_ALIGNMENT
            from docx.oxml.ns import qn

            doc = Document()

            # Base font
            style = doc.styles['Normal']
            font = style.font
            font.name = 'Calibri'
            font.size = Pt(11)
            font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

            # Title
            title_para = doc.add_heading(title, level=0)
            title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in title_para.runs:
                run.font.size = Pt(24)
                run.font.color.rgb = RGBColor(0x1a, 0x3a, 0x5c)

            # Author + date
            import datetime
            meta = doc.add_paragraph()
            meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = meta.add_run(f"Author: {author}  |  {datetime.date.today().strftime('%B %d, %Y')}")
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            doc.add_paragraph()

            # Content sections
            if content_sections:
                for section in content_sections:
                    stype = section.get("type", "paragraph")
                    text = section.get("text", "")
                    if stype == "heading":
                        level = min(section.get("level", 1), 4)
                        h = doc.add_heading(text, level=level)
                        for run in h.runs:
                            run.font.color.rgb = RGBColor(0x1a, 0x3a, 0x5c)
                    elif stype == "bullet":
                        doc.add_paragraph(text, style="List Bullet")
                    elif stype == "table":
                        headers = section.get("headers", [])
                        rows = section.get("rows", [])
                        if headers:
                            table = doc.add_table(rows=1 + len(rows), cols=len(headers))
                            table.style = "Table Grid"
                            hdr_row = table.rows[0]
                            for i, h in enumerate(headers):
                                cell = hdr_row.cells[i]
                                cell.text = h
                                for p in cell.paragraphs:
                                    for run in p.runs:
                                        run.bold = True
                                        run.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
                            for ri, row_data in enumerate(rows):
                                for ci, val in enumerate(row_data):
                                    table.rows[ri + 1].cells[ci].text = str(val)
                    else:
                        doc.add_paragraph(text)
            else:
                doc.add_paragraph("Document created by TOM.")

            # Save
            out_path = os.path.join(self.output_dir, file_name)
            doc.save(out_path)
            return {"status": "success", "message": f"Word document saved: {out_path}", "path": out_path}
        except ImportError:
            return {"status": "error", "message": "python-docx not installed. Run: pip install python-docx"}
        except Exception as e:
            return {"status": "error", "message": f"Word document creation failed: {str(e)}"}

    async def create_excel_spreadsheet(self, file_name: str, title: str = "Spreadsheet",
                                        data: Optional[List[List]] = None) -> Dict[str, Any]:
        """Create an Excel spreadsheet with data."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter

            wb = Workbook()
            ws = wb.active
            ws.title = title[:31]

            header_fill = PatternFill(start_color="1A3A5C", end_color="1A3A5C", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=11)
            border_side = Side(style="thin", color="CCCCCC")
            cell_border = Border(left=border_side, right=border_side,
                                  top=border_side, bottom=border_side)

            if data:
                for ri, row in enumerate(data, 1):
                    for ci, val in enumerate(row, 1):
                        cell = ws.cell(row=ri, column=ci, value=val)
                        cell.border = cell_border
                        if ri == 1:
                            cell.font = header_font
                            cell.fill = header_fill
                            cell.alignment = Alignment(horizontal="center")
                # Auto-fit columns
                for col in ws.columns:
                    max_len = max((len(str(c.value)) for c in col if c.value), default=10)
                    ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)

            out_path = os.path.join(self.output_dir, file_name)
            wb.save(out_path)
            return {"status": "success", "message": f"Excel saved: {out_path}", "path": out_path}
        except ImportError:
            return {"status": "error", "message": "openpyxl not installed. Run: pip install openpyxl"}
        except Exception as e:
            return {"status": "error", "message": f"Excel creation failed: {str(e)}"}

    # ── Bridge methods to document_creator ────────────────────────────────

    async def create_excel_workbook(self, file_name: str, sheets=None) -> Dict[str, Any]:
        try:
            from tools.document_creator import create_excel_spreadsheet as _create
            data = None
            if sheets and len(sheets) > 0:
                sheet = sheets[0]
                headers = sheet.get("headers", [])
                rows = sheet.get("data", [])
                if headers and rows:
                    data = [dict(zip(headers, row)) for row in rows]
            result = _create(
                title=file_name.replace(".xlsx", "").replace("_", " ").title(),
                data=data,
                filename=file_name,
            )
            return result
        except Exception as e:
            return await self.create_excel_spreadsheet(file_name)

    async def create_powerpoint(self, file_name: str, slides=None,
                                 title: str = "Presentation",
                                 theme: dict = None) -> Dict[str, Any]:
        try:
            from tools.document_creator import create_presentation as _create
            slides_content = None
            if slides:
                slides_content = []
                for s in slides:
                    entry = {
                        "title": s.get("title", ""),
                        "bullets": s.get("content", []),
                        "notes": s.get("notes", s.get("subtitle", "")),
                        "layout": s.get("layout", "content"),
                    }
                    if s.get("table_data"):
                        entry["table_data"] = s["table_data"]
                    if s.get("chart_data"):
                        entry["chart_data"] = s["chart_data"]
                    slides_content.append(entry)
            result = _create(
                title=title,
                slides_content=slides_content,
                filename=file_name,
                theme=theme,
            )
            return result
        except Exception as e:
            return {"status": "error", "message": f"Presentation creation failed: {e}"}

    async def create_pdf_report(self, file_name: str, title: str = "Report",
                                 author: str = "TOM AI", sections=None) -> Dict[str, Any]:
        try:
            from tools.document_creator import create_pdf_report as _create
            content_lines = []
            if sections:
                for sec in sections:
                    sec_type = sec.get("type", "paragraph")
                    text = sec.get("text", "")
                    level = sec.get("level", 1)
                    if sec_type == "heading":
                        content_lines.append(f"{'#' * level} {text}")
                    elif sec_type == "bullet":
                        content_lines.append(f"- {text}")
                    elif sec_type == "table":
                        headers = sec.get("headers", [])
                        rows = sec.get("data", [])
                        if headers:
                            content_lines.append(" | ".join(headers))
                            for row in rows:
                                content_lines.append(" | ".join(str(v) for v in row))
                    else:
                        content_lines.append(text)
            content = "\n".join(content_lines) if content_lines else f"Report about {title}"
            result = _create(
                title=title,
                content=content,
                author=author,
                filename=file_name,
            )
            return result
        except Exception as e:
            return {"status": "error", "message": f"PDF creation failed: {e}"}
